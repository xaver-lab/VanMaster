"""Einzelteile: was gebaut und zugeschnitten wird, nicht was gekauft wird.

Ein Holzbrett 1200x400x15 gehört nicht in die Stückliste — man kauft keine
Bretter, man kauft eine Platte. Es gehört aber in den Arbeitsbereich, mit Maß
und Material. Über `teil_id` kann ein Einzelteil auf den Stücklisten-Artikel
zeigen, aus dem es entsteht; leer ist ausdrücklich erlaubt.

data/bauteile.csv ist die Wahrheit, Excel ist eine Ausleihe — wie bei parts.py.
Lesen über ``tools.kern.lesen.einzelteile_lesen``, Feld ändern und Anlegen
über ``tools.kern.tabellen`` (Quelle ``claude``); der Excel-Import bleibt ein
Massenschreiben hier, aber über ``tools.kern.datei``.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from .common import (
    BAUTEILE_CSV, BAUTEILE_XLSX, BAUTEIL_ART, BAUTEIL_STATUS, MASSQUELLE,
    fail, table,
)
from .kern import datei as kern_datei
from .kern import tabellen as kern_tabellen
from .kern.format import (
    EINZELTEIL_BERECHNET as COMPUTED, EINZELTEIL_FELDER as FIELDS,
)
from .kern.lesen import einzelteile_lesen

# FIELDS: Spaltenreihenfolge von data/bauteile.csv. COMPUTED: nur in der
# Excel-Ausleihe, beim Import ignoriert. Beide aus tools/kern/format.py.

ZAHLFELDER = ("laenge_mm", "breite_mm", "dicke_mm", "anzahl")


def load() -> list[dict]:
    return [{f: getattr(r, f) for f in FIELDS} for r in einzelteile_lesen()]


def save(rows: list[dict]) -> None:
    """Schreibt die ganze Datei neu — für den Excel-Import. Über
    ``kern.datei``, damit Hash-Versionsschutz und atomares Schreiben auch
    hier gelten."""
    BAUTEILE_CSV.parent.mkdir(parents=True, exist_ok=True)
    rows = sorted(rows, key=lambda r: (r.get("bereich", ""), r.get("titel", "")))
    puffer = io.StringIO()
    writer = csv.DictWriter(puffer, fieldnames=FIELDS, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({f: row.get(f, "") for f in FIELDS})
    kern_datei.schreiben(BAUTEILE_CSV, puffer.getvalue(), None)


def num(value, default: float = 0.0) -> float:
    try:
        return float(str(value).replace(",", ".").strip() or default)
    except ValueError:
        return default


def anzahl(row: dict) -> float:
    return num(row.get("anzahl"), 1)


def flaeche(row: dict) -> float:
    """Fläche aller Stücke in m² — 0.0, solange Länge oder Breite fehlt."""
    l, b = num(row.get("laenge_mm")), num(row.get("breite_mm"))
    if not l or not b:
        return 0.0
    return l * b * anzahl(row) / 1_000_000


def laufmeter(row: dict) -> float:
    """Gesamtlänge in m — für Leisten, Rohre und Kabel."""
    return num(row.get("laenge_mm")) * anzahl(row) / 1000


def flaeche_summe(rows: list[dict]) -> float:
    return sum(flaeche(r) for r in rows)


def mass_text(row: dict) -> str:
    teile = [f"{num(row[f]):g}" for f in ("laenge_mm", "breite_mm", "dicke_mm")
             if num(row.get(f))]
    return " x ".join(teile) + (" mm" if teile else "")


def new_id(titel: str, rows: list[dict]) -> str:
    from .common import slug
    base = slug(titel) or "bauteil"
    existing = {r.get("id", "") for r in rows}
    if base not in existing:
        return base
    i = 2
    while f"{base}-{i}" in existing:
        i += 1
    return f"{base}-{i}"


def find(bauteil_id: str, rows: list[dict] | None = None) -> dict | None:
    rows = load() if rows is None else rows
    bauteil_id = bauteil_id.strip().lower()
    for row in rows:
        if row["id"].lower() == bauteil_id:
            return row
    return None


def nach_bereich(rows: list[dict]) -> dict[str, list[dict]]:
    gruppen: dict[str, list[dict]] = {}
    for r in rows:
        gruppen.setdefault(r["bereich"], []).append(r)
    return gruppen


# ---------------------------------------------------------------- Abfragen

def filtered(rows: list[dict], **f) -> list[dict]:
    def keep(row: dict) -> bool:
        for key, value in f.items():
            if not value:
                continue
            if key == "text":
                if str(value).lower() not in " ".join(row.values()).lower():
                    return False
            elif row.get(key, "").lower() != str(value).lower():
                return False
        return True
    return [r for r in rows if keep(r)]


def query_text(rows: list[dict]) -> str:
    if not rows:
        return "Keine Einzelteile gefunden."
    body = [[
        r["id"], r["titel"][:32], r["bereich"], r["art"],
        mass_text(r), f"{anzahl(r):g}", r["material"][:22],
        r["massquelle"], r["status"],
    ] for r in rows]
    head = ["id", "titel", "bereich", "art", "maß", "anz",
            "material", "quelle", "status"]
    foot = f"\n{len(rows)} Einzelteile"
    qm = flaeche_summe(rows)
    if qm:
        foot += f" · {qm:.2f} m² Fläche"
    return table(body, head) + foot


def overview_text() -> str:
    rows = load()
    if not rows:
        return ("Noch keine Einzelteile. Anlegen mit "
                "`camper bauteile add --titel ... --bereich ...`")
    lines = []
    for bereich, brows in sorted(nach_bereich(rows).items()):
        verbaut = len([r for r in brows if r["status"] == "Verbaut"])
        qm = flaeche_summe(brows)
        lines.append([bereich, str(len(brows)), f"{verbaut}",
                      f"{qm:.2f} m²" if qm else "—"])
    body = table(lines, ["bereich", "teile", "verbaut", "fläche"])
    ohne = len([r for r in rows if not mass_text(r)])
    return (f"{body}\n\nGesamt: {len(rows)} Einzelteile · "
            f"{flaeche_summe(rows):.2f} m² · {ohne} ohne Maß")


def set_field(bauteil_id: str, field: str, value: str) -> str:
    if field not in FIELDS or field == "id":
        fail(f"Unbekanntes Feld '{field}'. Erlaubt: {', '.join(FIELDS[1:])}")
    row = find(bauteil_id)
    if row is None:
        fail(f"Kein Einzelteil mit der Kennung '{bauteil_id}'.")
    if field == "status" and value not in BAUTEIL_STATUS:
        fail(f"Status muss einer von {', '.join(BAUTEIL_STATUS)} sein.")
    if field == "art" and value not in BAUTEIL_ART:
        fail(f"Art muss eine von {', '.join(BAUTEIL_ART)} sein.")
    if field == "massquelle" and value not in MASSQUELLE:
        fail(f"Maßquelle muss eine von {', '.join(MASSQUELLE)} sein.")
    alt = row[field]
    try:
        kern_tabellen.einzelteil_feld_setzen(row["id"], field, value, None,
                                             quelle="claude")
    except kern_tabellen.Ungueltig as fehler:
        fail(str(fehler))
    return f"{row['titel']}: {field} {alt or '—'} → {value}"


def add(titel: str, bereich: str, **extra) -> str:
    felder = {"titel": titel, "bereich": bereich}
    felder.update({k: str(v) for k, v in extra.items() if v})
    try:
        neue_id, _ = kern_tabellen.einzelteil_anlegen(felder, None, quelle="claude")
    except kern_tabellen.Ungueltig as fehler:
        fail(str(fehler))
    row = find(neue_id)
    mass = mass_text(row) if row else ""
    return (f"Aufgenommen: {titel} ({neue_id}, {bereich})"
            + (f" · {mass}" if mass else " · ohne Maß"))


# ---------------------------------------------------------------------- Excel

SPALTENBREITEN = {"id": 24, "titel": 34, "material": 26, "notiz": 30,
                  "teil_id": 20, "fuer_aufgabe": 20, "massquelle": 12}


def to_excel(path: Path | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path = path or BAUTEILE_XLSX
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = sorted(load(), key=lambda r: (r["bereich"], r["titel"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Einzelteile"
    header = FIELDS + COMPUTED
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="333A45")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "C2"

    l_col = get_column_letter(FIELDS.index("laenge_mm") + 1)
    b_col = get_column_letter(FIELDS.index("breite_mm") + 1)
    a_col = get_column_letter(FIELDS.index("anzahl") + 1)
    for i, row in enumerate(rows, start=2):
        ws.append([row.get(f, "") for f in FIELDS])
        ws.cell(row=i, column=len(FIELDS) + 1).value = (
            f"=IFERROR(N({l_col}{i})*N({b_col}{i})*N({a_col}{i})/1000000,0)"
        )
        ws.cell(row=i, column=len(FIELDS) + 2).value = (
            f"=IFERROR(N({l_col}{i})*N({a_col}{i})/1000,0)"
        )

    letzte = max(ws.max_row, 2)
    for spalte, werte in (("art", BAUTEIL_ART), ("status", BAUTEIL_STATUS),
                          ("massquelle", MASSQUELLE)):
        col = get_column_letter(FIELDS.index(spalte) + 1)
        dv = DataValidation(type="list",
                            formula1='"' + ",".join(werte) + '"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{letzte + 200}")

    for i, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = SPALTENBREITEN.get(name, 13)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{letzte}"

    wb.save(path)
    return f"{len(rows)} Einzelteile → {path}"


def read_excel(path: Path | None = None) -> list[dict]:
    from openpyxl import load_workbook

    path = path or BAUTEILE_XLSX
    if not path.exists():
        fail(f"Keine Excel-Datei unter {path}. Erst 'bauteile excel' laufen lassen.")
    wb = load_workbook(path, data_only=False)
    ws = wb["Einzelteile"] if "Einzelteile" in wb.sheetnames else wb.active
    header = [str(c.value or "").strip() for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        raw = {h: values[i] for i, h in enumerate(header) if h in FIELDS}
        if not str(raw.get("titel") or "").strip():
            continue
        rows.append({f: ("" if raw.get(f) is None else str(raw.get(f, "")).strip())
                     for f in FIELDS})
    return rows


def diff(neu: list[dict], alt: list[dict] | None = None) -> dict:
    alt = load() if alt is None else alt
    alt_map = {r["id"]: r for r in alt if r["id"]}
    neu_map: dict[str, dict] = {}
    for row in neu:
        if not row["id"]:
            row["id"] = new_id(row["titel"], list(alt) + list(neu_map.values()))
        neu_map[row["id"]] = row
    geaendert = []
    for bid, row in neu_map.items():
        vorher = alt_map.get(bid)
        if not vorher:
            continue
        felder = [f for f in FIELDS if (vorher.get(f) or "") != (row.get(f) or "")]
        if felder:
            geaendert.append((row["titel"], felder, vorher, row))
    return {
        "neu": [r for bid, r in neu_map.items() if bid not in alt_map],
        "entfernt": [r for bid, r in alt_map.items() if bid not in neu_map],
        "geaendert": geaendert,
        "rows": list(neu_map.values()),
    }


def diff_text(d: dict) -> str:
    lines = []
    for row in d["neu"]:
        lines.append(f"  + {row['titel']} ({row['bereich']}, {mass_text(row)})")
    for row in d["entfernt"]:
        lines.append(f"  - {row['titel']} ({row['bereich']})")
    for titel, felder, vorher, nachher in d["geaendert"]:
        aend = ", ".join(f"{f}: {vorher.get(f) or '—'} → {nachher.get(f) or '—'}"
                         for f in felder)
        lines.append(f"  ~ {titel}: {aend}")
    if not lines:
        return "Keine Unterschiede — die Excel-Datei entspricht der CSV."
    kopf = (f"{len(d['geaendert'])} geändert, {len(d['neu'])} neu, "
            f"{len(d['entfernt'])} entfernt")
    return kopf + "\n" + "\n".join(lines)


def import_excel(path: Path | None = None, apply: bool = False) -> str:
    d = diff(read_excel(path))
    text = diff_text(d)
    if not d["neu"] and not d["entfernt"] and not d["geaendert"]:
        return text
    if not apply:
        return text + "\n\nNichts geschrieben. Mit --apply übernehmen."
    save(d["rows"])
    return text + "\n\nÜbernommen nach data/bauteile.csv."
